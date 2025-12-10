# -*- coding: utf-8 -*-
"""
전략 로드맵 수립 노드 모듈
PDF 파일에서 전략 내용을 추출하고, 웹 검색을 통해 3개년 로드맵을 수립
"""
import os
import sys
import io
import tempfile
from typing import TypedDict, List
from langgraph.graph import StateGraph, END, START
from langgraph.prebuilt import ToolNode
from langchain.tools import tool
from docx import Document
from docx.oxml.ns import qn

# LangChain 관련 import
from langchain_core.messages import HumanMessage, BaseMessage
from langchain_openai import ChatOpenAI
from langchain_core.runnables import RunnableLambda 

# Web Search
from langchain_tavily import TavilySearch

# Vector DB
from langchain_community.vectorstores import FAISS
from langchain_openai import OpenAIEmbeddings
from langchain_community.document_loaders import PyPDFLoader
from langchain_classic.text_splitter import RecursiveCharacterTextSplitter

# Windows 환경에서 한글 출력을 위한 인코딩 설정
if sys.platform == 'win32':
    os.environ['PYTHONIOENCODING'] = 'utf-8'
    try:
        if hasattr(sys.stdout, 'reconfigure'):
            sys.stdout.reconfigure(encoding='utf-8', errors='replace')
            sys.stderr.reconfigure(encoding='utf-8', errors='replace')
    except:
        pass


@tool
def save_to_word(content: str, filename: str = "output.docx") -> str:
    """
    한글이 깨지지 않도록 폰트를 설정해 Word 파일(.docx)로 저장하는 LangChain Tool.
    Args:
        content (str): 저장할 텍스트 내용
        filename (str): 저장할 파일 이름
    Returns:
        str: 저장 성공/실패 메시지
    """
    try:
        doc = Document()
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(content)

        # 🔥 한글 폰트 설정 부분 (매우 중요)
        run.font.name = "NanumGothic"            # 문단 기본 글꼴
        run._element.rPr.rFonts.set(qn('w:eastAsia'), "NanumGothic")  # 한글 폰트 설정

        doc.save(filename)
        return f"📄 Word 파일 저장 완료: {filename}"
    except Exception as e:
        return f"❌ Word 저장 중 오류 발생: {str(e)}"


# State는 LangGraph에서 노드 간 데이터를 전달하는 구조입니다.
class State(TypedDict):
    subject: str                # 주제
    messages: List[BaseMessage]  # 질문
    plan: str                   # 계획
    web_Industry: str           # 산업 데이터
    web_Technology: str        # 기술 데이터
    web_Policy: str            # 정부정책 데이터
    web_Regulation: str         # 규제/법령 데이터
    answer: str                # 최종 생성 답변
    search: str                # 검색용 전략 원문


def create_strategy_plan_chain(pdf_file_path: str, api_key: str = None, tavily_api_key: str = None, 
                                save_word_file: bool = False, output_dir: str = None, verbose: bool = False):
    """
    전략 로드맵 수립 체인 생성 함수
    
    Args:
        pdf_file_path (str): 전략 PDF 파일 경로
        api_key (str, optional): OpenAI API 키. None이면 환경 변수에서 가져옴
        tavily_api_key (str, optional): Tavily API 키. None이면 환경 변수에서 가져옴
        save_word_file (bool): Word 파일로 저장할지 여부
        output_dir (str, optional): Word 파일 저장 디렉토리
        verbose (bool): 상세 로그 출력 여부
        
    Returns:
        chain: 전략 로드맵 수립 체인 객체
    """
    try:
        # API 키 설정
        if api_key:
            os.environ["OPENAI_API_KEY"] = api_key
        if tavily_api_key:
            os.environ["TAVILY_API_KEY"] = tavily_api_key
        
        # API 키 확인
        if not os.environ.get("OPENAI_API_KEY"):
            raise ValueError("OpenAI API Key가 설정되지 않았습니다.")
        if not os.environ.get("TAVILY_API_KEY"):
            raise ValueError("Tavily API Key가 설정되지 않았습니다.")
        
        # LLM 초기화
        llm = ChatOpenAI(model="gpt-5-mini")
        
        # PDF 파일 로드
        if not os.path.exists(pdf_file_path):
            raise FileNotFoundError(f"PDF 파일을 찾을 수 없습니다: {pdf_file_path}")
        
        loader = PyPDFLoader(pdf_file_path)
        docs = loader.load()
        
        if verbose:
            print(f"✅ PDF 파일 로드 완료: {len(docs)}개 페이지")
        
        # PDF 전체 텍스트 로드
        full_text = "\n".join([page.page_content for page in docs])
        
        # LLM을 이용해 '세부전략' 단위로 내용 발췌
        extraction_prompt = f"""
다음 문서는 사업 전략에 대한 내용입니다.
문서 내용 중 '세부전략 1', '세부전략 2', '세부전략 3' 등 각 전략 항목의 타이틀과 세부 내용을 **원문 그대로** 추출해서 정리해주세요.
요약하지 말고, 있는 그대로의 텍스트를 가져와야 합니다.

[문서 내용]
{full_text}

[출력 형식]
1. [세부전략 1 제목]
- 내용...
2. [세부전략 2 제목]
- 내용...
...
"""
        
        strategy_content_msg = llm.invoke(extraction_prompt)
        strategy_content = strategy_content_msg.content
        
        if verbose:
            print("✅ 전략 내용 추출 완료")
        
        # Web Search Tool 초기화
        tavily_Industry = TavilySearch(max_results=5)
        tavily_Technology = TavilySearch(max_results=5)
        tavily_Policy = TavilySearch(max_results=5)
        tavily_Regulation = TavilySearch(max_results=5)
        
        # Manager 1: Planner (작업 배분)
        def manager_planner(state: State):
            """프로젝트를 분석하고 작업을 배분하는 Manager"""
            subject = state["subject"]
            
            if verbose:
                print(f"=== Manager (Planner) 실행 ===")
            
            prompt = f"""
당신은 전략 프로젝트 PM 입니다.
다음은 우리 회사의 핵심 [사업 전략 목록] 입니다. 

[사업 전략 목록]
{subject}

이 전략을 실행하기 위해, 향후 3개년 로드맵을 수립하려고 합니다. 
각 검색 담당자(Industry, Technology, Policy, Regulation)가 위의 '세부전략1, 2, 3, ...'과 
관련된 어떤 정보를 찾아야 하는지 구체적인 지시사항을 작성하세요.

예시:
- Web_Industry : 벽돌/건자재 산업의 품질 관리 트렌드 및 경쟁사 동향 검색
- Web_Technology : AI 기반 제조 공정 품질 예측(Feed-Forward) 기술 및 데이터 통합 플랫폼 사례 검색
...
"""
            
            response = llm.invoke([HumanMessage(content=prompt)])
            work_plan = response.content
            
            if verbose:
                print(f"작업 배분 계획 생성 완료")
            
            return {"plan": work_plan, "messages": [response]}
        
        def web_Industry(state: State):
            plan = state["plan"]
            strategies = state["search"]
            
            if verbose:
                print("=== 산업 동향 검색 ===")
            
            query_gen_prompt = f"""
[전체 계획]: {plan}
[상세 계획]: {strategies}

위 전략들을 실현하기 위해 가장 시급하게 조사해야 할 '산업 동향' 검색어를 선정해줘
(예: '벽돌/건자재 산업의 품질 관리 트렌드 및 경쟁사 동향 검색')
"""
            
            query_gen_response = llm.invoke(query_gen_prompt)
            search_keyword = query_gen_response.content
            
            if verbose:
                print(f"생성된 검색어: {search_keyword}")
            
            search_result = tavily_Industry.invoke(search_keyword)
            
            system_message = f"""
[상세 전략]: {strategies}
[검색된 산업 정보]: {search_result}

검색된 산업동향 정보를 바탕으로, 우리 전략(세부전략1, 세부전략2, 세부전략3, ...)에 참고해야 하는 
산업동향을 정리해줘.
"""
            
            rag_chain_response = llm.invoke(system_message)
            return {"web_Industry": rag_chain_response.content}
        
        def web_Technology(state: State):
            plan = state["plan"]
            strategies = state["search"]
            
            if verbose:
                print("=== 기술 동향 검색 ===")
            
            query_gen_prompt = f"""
[전체 계획]: {plan}
[상세 계획]: {strategies}

위 전략들을 실현하기 위해 가장 시급하게 조사해야 할 '최신 기술 트렌드' 검색어를 선정해줘
(예: 'AI Quality Control in Brick Manufacturing' 또는 'Smart Factory Feed-Forward control')
"""
            
            query_gen_response = llm.invoke(query_gen_prompt)
            search_keyword = query_gen_response.content
            
            if verbose:
                print(f"생성된 검색어: {search_keyword}")
            
            search_result = tavily_Technology.invoke(search_keyword)
            
            system_message = f"""
[상세 전략]: {strategies}
[검색된 기술 정보]: {search_result}

검색된 기술 정보를 바탕으로, 우리 전략(세부전략1, 세부전략2, 세부전략3, ...)에 적용 가능한 
기술적 시사점과 도입 방향을 정리해줘.
"""
            
            rag_chain_response = llm.invoke(system_message)
            return {"web_Technology": rag_chain_response.content}
        
        def web_Policy(state: State):
            plan = state["plan"]
            strategies = state["search"]
            
            if verbose:
                print("=== 정책 동향 검색 ===")
            
            query_gen_prompt = f"""
[전체 계획]: {plan}
[상세 계획]: {strategies}

위 전략들을 실현하기 위해 가장 시급하게 조사해야 할 '정책 동향' 검색어를 선정해줘
"""
            
            query_gen_response = llm.invoke(query_gen_prompt)
            search_keyword = query_gen_response.content
            
            if verbose:
                print(f"생성된 검색어: {search_keyword}")
            
            search_result = tavily_Policy.invoke(search_keyword)
            
            system_message = f"""
[상세 전략]: {strategies}
[검색된 정책 정보]: {search_result}

검색된 정책 정보를 바탕으로, 우리 전략(세부전략1, 세부전략2, 세부전략3, ...)에 참고해야 하는 
정책 동향에 대해 정리해줘.
"""
            
            rag_chain_response = llm.invoke(system_message)
            return {"web_Policy": rag_chain_response.content}
        
        def web_Regulation(state: State):
            plan = state["plan"]
            strategies = state["search"]
            
            if verbose:
                print("=== 규제/법령 동향 검색 ===")
            
            query_gen_prompt = f"""
[전체 계획]: {plan}
[상세 계획]: {strategies}

위 전략들을 실현하기 위해 가장 시급하게 조사해야 할 '규제/법령 동향' 검색어를 선정해줘
"""
            
            query_gen_response = llm.invoke(query_gen_prompt)
            search_keyword = query_gen_response.content
            
            if verbose:
                print(f"생성된 검색어: {search_keyword}")
            
            search_result = tavily_Regulation.invoke(search_keyword)
            
            system_message = f"""
[상세 전략]: {strategies}
[검색된 규제/법령 정보]: {search_result}

검색된 규제/법령 정보를 바탕으로, 우리 전략(세부전략1, 세부전략2, 세부전략3, ...)에 참고해야 하는 
규제/법령 동향에 대해 정리해줘.
"""
            
            rag_chain_response = llm.invoke(system_message)
            return {"web_Regulation": rag_chain_response.content}
        
        def integrator(state: State):
            """Workers의 결과를 취합하는 Manager"""
            subject = state["subject"]
            web_Industry = state["web_Industry"]
            web_Technology = state["web_Technology"]
            web_Policy = state["web_Policy"]
            web_Regulation = state["web_Regulation"]
            
            if verbose:
                print(f"=== Manager (Integrator) 실행 ===")
            
            prompt = f"""
주제:
{subject}

산업동향 결과:
{web_Industry}

기술동향 결과:
{web_Technology}

정책동향 결과:
{web_Policy}

규제/법령 결과:
{web_Regulation}

위 내용들을 통합 및 참고하여
입력된 세부전략별에 대해 아래 양식으로 향후 3개년 전략 로드맵을 반기 단위로 만들어줘

1.+1년
- 상반기 :
- 하반기 :
2.+2년
- 상반기 :
- 하반기 :
3.+3년
- 상반기 :
- 하반기 :
"""
            
            response = llm.invoke([HumanMessage(content=prompt)])
            final_design = response.content
            
            if verbose:
                print(f"통합 완료")
            
            return {"answer": final_design, "messages": [response]}
        
        def tool_node(state: State):
            """Word 파일 저장을 위한 툴 노드"""
            messages = state["messages"][-1].content if state["messages"] else ""
            answer = state["answer"]
            
            if not save_word_file:
                return {"answer": answer}
            
            user_query = f"""
{messages}

내용 : {answer}
"""
            
            tools = [save_to_word]
            llm_with_tools = llm.bind_tools(tools)
            response = llm_with_tools.invoke(user_query)
            
            tool_node_instance = ToolNode(tools)
            
            # 출력 디렉토리 설정
            if output_dir:
                os.makedirs(output_dir, exist_ok=True)
                filename = os.path.join(output_dir, "3개년_전략_로드맵.docx")
            else:
                filename = "3개년_전략_로드맵.docx"
            
            # 파일명을 툴에 전달하기 위해 메시지 수정
            if hasattr(response, 'tool_calls') and response.tool_calls:
                for tool_call in response.tool_calls:
                    if tool_call.get('name') == 'save_to_word':
                        tool_call['args']['filename'] = filename
            
            result = tool_node_instance.invoke({"messages": [response]})
            
            return {"answer": answer}
        
        # LangGraph 워크플로우 구성
        workflow = StateGraph(State)
        
        workflow.add_node("plan", manager_planner)
        workflow.add_node("web_Industry", web_Industry)
        workflow.add_node("web_Technology", web_Technology)
        workflow.add_node("web_Policy", web_Policy)
        workflow.add_node("web_Regulation", web_Regulation)
        workflow.add_node("integrator", integrator)
        workflow.add_node("word_file", tool_node)
        
        # 엣지 연결
        workflow.add_edge(START, "plan")
        workflow.add_edge("plan", "web_Industry")
        workflow.add_edge("plan", "web_Technology")
        workflow.add_edge("plan", "web_Policy")
        workflow.add_edge("plan", "web_Regulation")
        workflow.add_edge("web_Industry", "integrator")
        workflow.add_edge("web_Technology", "integrator")
        workflow.add_edge("web_Policy", "integrator")
        workflow.add_edge("web_Regulation", "integrator")
        workflow.add_edge("integrator", "word_file")
        workflow.add_edge("word_file", END)
        
        # 그래프 컴파일
        app = workflow.compile()
        
        # 체인 래퍼 함수
        def strategy_plan_chain(inputs: dict) -> str:
            """전략 로드맵 수립 체인 실행"""
            question = inputs.get("question", "3개년 전략 로드맵 수립해줘")
            
            result = app.invoke({
                "messages": [HumanMessage(content=question)],
                "subject": strategy_content,
                "search": strategy_content
            })
            
            return result.get("answer", "로드맵 생성에 실패했습니다.")
        
        if verbose:
            print("✅ 전략 로드맵 수립 체인 구성 완료")
        
        return RunnableLambda(strategy_plan_chain)
        
    except Exception as e:
        raise Exception(f"전략 로드맵 노드 초기화 실패: {str(e)}")


# 직접 실행 시 (기존 코드 호환성 유지)
if __name__ == "__main__":
    from dotenv import load_dotenv
    
    env_path = r"C:\Users\user\Desktop\PJT1\.env"
    load_dotenv(dotenv_path=env_path)
    
    pdf_path = r"C:\Users\user\Desktop\PJT1\사업전략\벽돌사업전략.pdf"
    
    chain = create_strategy_plan_chain(
        pdf_file_path=pdf_path,
        save_word_file=True,
        verbose=True
    )
    
    result = chain.invoke({"question": "3개년 전략 로드맵 수립해줘"})
    print("\n" + "="*60)
    print("📊 전략 로드맵 수립 결과")
    print("="*60)
    print(result)
    print("="*60)

